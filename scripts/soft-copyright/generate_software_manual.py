"""生成软著配套的软件使用说明书 DOCX 与 PDF。"""

from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image as PdfImage,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = ROOT / "docs" / "soft-copyright" / "deliverables"
SCREENSHOT_DIR = DEFAULT_OUTPUT / "screenshots"
SOFTWARE_NAME = "新材料数据分析---拓扑感知多核集成聚类分析系统"
SOFTWARE_SHORT_NAME = "新材料数据分析系统"
SOFTWARE_VERSION = "V1.0"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5C6873"
LIGHT_FILL = "E8EEF5"


SECTIONS: list[tuple[str, list[str]]] = [
    (
        "1. 软件概述",
        [
            "本系统面向新材料数据分析场景，围绕基础聚类结果管理、拓扑感知多核集成聚类任务、运行监控、性能评估、可视化与结果交付形成统一的桌面端 Web 工作台。",
            "系统采用前后端分离架构。前端负责认证交互、数据管理、任务中心和结果展示；后端负责权限校验、MAT 文件安全解析、任务调度、结果持久化、导出与审计。算法计算在独立子进程中执行，避免阻塞 Web 服务。",
            "本说明书对应软件版本 V1.0，页面、接口名称和源码清单均以当前仓库运行时代码为准。",
        ],
    ),
    (
        "2. 运行环境与系统组成",
        [
            "客户端使用桌面浏览器访问，不面向移动端适配。建议分辨率不低于 1366×768，浏览器启用 JavaScript 和 Cookie。",
            "前端使用 React、TypeScript、Vite 与 ECharts；后端使用 FastAPI、SQLAlchemy、NumPy、SciPy 和 scikit-learn；业务数据保存到 MySQL，数据集、结果和导出文件保存到受控存储目录。",
            "开发环境默认前端地址为 http://127.0.0.1:5173，后端健康检查地址为 http://127.0.0.1:8000/api/health。生产部署应通过反向代理提供 HTTPS，并配置独立的环境变量和数据库迁移步骤。",
        ],
    ),
    (
        "3. 用户认证与会话",
        [
            "用户从首页进入登录或注册页面。注册流程可启用邮箱验证码与 Cloudflare Turnstile，人机验证是否开启由部署配置决定。",
            "登录成功后，access token 只保存在浏览器内存中；refresh token 由后端写入 HttpOnly Cookie。access token 过期时，统一请求层只发起一次并发共享刷新，并在刷新成功后重试原请求一次。",
            "refresh token 每次使用后轮换，旧令牌立即撤销；检测到旧令牌重复使用时撤销该用户全部会话。退出登录会撤销数据库会话并清除 Cookie。",
        ],
    ),
    (
        "4. 数据管理",
        [
            "数据管理页面提供上传、追加、替换、重命名、版本查看、质量复核、导出和删除操作。系统只接受允许的 MAT 文件扩展名，并在写入正式目录前完成大小、变量数量、矩阵维度和数值结构检查。",
            "文件以分块方式写入临时文件，MAT 解析在受控子进程中执行并设置超时。系统拒绝 NaN、Inf、对象数组和异常结构，同时检查单用户存储配额。数据库提交或文件写入任一步失败时，临时文件和数据库状态都会回滚。",
            "数据集关联任务数量统计 analysis_tasks 的全部状态，最近分析时间取最近一条成功任务的 completed_at。只要存在任意正式分析任务引用，系统就禁止物理删除数据集，并提示先处理关联任务。",
        ],
    ),
    (
        "5. 任务中心与执行流程",
        [
            "用户选择通过质量检查的数据集，设置算法模式、基础聚类数量、核参数、运行次数和迭代上限后创建任务。草稿任务可编辑或删除，排队任务由执行器领取，运行任务持续更新心跳和进度。",
            "执行器在数据库事务内原子领取 queued 任务，并写入 worker_id、heartbeat_at 和 retry_count，避免多个服务实例重复领取。任务超过最大运行时间时，系统终止完整子进程树并记录失败原因。",
            "子进程只接收经过归一化的任务参数和受控数据文件路径。计算完成后生成 manifest、标签、矩阵、指标和 CSV 等产物；执行器在锁定任务记录后一次性持久化结果，已成功结果不会被迟到进程覆盖。",
        ],
    ),
    (
        "6. 结果、导出与日志",
        [
            "结果页面统一读取后端任务结果 envelope，根据任务状态显示草稿、排队、运行、失败、取消或已完成视图。已完成任务可查看 CA 协关联矩阵、核函数配置、多核学习、性能指标、可视化、综合分析和报告。",
            "页面不保存固定指标、矩阵、散点或模拟日志。所有任务名称、运行进度、指标、矩阵、标签、产物列表和导出记录均来自后端真实任务数据。",
            "数据集导出由后端读取真实文件并生成 ZIP；任务导出在受控临时目录生成，数据库记录提交成功后才作为可下载产物公开。操作日志记录数据解析、任务状态转换、导出和异常，便于问题定位与审计。",
        ],
    ),
    (
        "7. 安全与可靠性边界",
        [
            "生产环境未配置至少 32 字节的 JWT 密钥时拒绝启动；示例 MAT 路由只在开发和测试环境注册，生产路由表不包含该接口。",
            "所有数据集和任务接口都执行当前用户权限校验。结果文件下载只允许访问任务清单登记且位于任务结果目录内的文件，防止路径穿越。",
            "数据库结构由 Alembic 管理，应用启动只核对迁移版本。数据库记录删除后，磁盘文件通过事务性清理队列重试，避免数据库与文件系统删除顺序不一致。",
        ],
    ),
    (
        "8. 源码材料范围",
        [
            "本次软著材料以原创平台代码为主体，包括 backend/main.py、backend/app/ 和 front/src/ 的运行时代码。测试、数据库迁移、SQL 快照、依赖、日志、构建产物、环境配置和二进制数据不进入源码 PDF 主体。",
            "backend/app/task_worker.py 保留在平台源码中，因为它承担任务参数转换、算法调用编排和结果产物标准化。ec_python_converted/ 算法转换实现暂不作为本次原创源码提交，相关样例只保存在 tests/fixtures/。",
            "源码清单、排除项、SHA-256 摘要、完整源码 PDF、60 页提交摘录和可追溯源码 ZIP 均由 source-manifest.json 和生成脚本自动产生。",
        ],
    ),
]


SCREENSHOTS = [
    ("01-首页.png", "图 1  系统首页：统一软件名称、版本和功能入口"),
    ("02-登录页.png", "图 2  用户登录：账号输入与人机验证区域"),
    ("03-分析工作台.png", "图 3  分析工作台：后端连接状态与无结果业务空状态"),
    ("04-数据管理.png", "图 4  数据管理：搜索、筛选、上传与空目录状态"),
    ("05-任务中心.png", "图 5  任务中心：统计、筛选、创建引导和任务列表"),
]


def _set_run_font(run, name: str = "Microsoft YaHei", size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _add_page_field(paragraph) -> None:
    paragraph.add_run("第 ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instruction, separate, text, end])
    paragraph.add_run(" 页")


def _configure_docx_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Number", "List Bullet"):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def _add_docx_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run(f"{SOFTWARE_SHORT_NAME} | 软件使用说明书")
    _set_run_font(run, size=8.5, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.paragraph_format.space_before = Pt(0)
    _add_page_field(footer_p)
    for footer_run in footer_p.runs:
        _set_run_font(footer_run, size=8.5, color=MUTED)


def _add_docx_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("软件著作权登记配套材料")
    _set_run_font(run, size=11, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run(SOFTWARE_NAME)
    _set_run_font(run, size=25, bold=True, color=DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(36)
    run = subtitle.add_run("软件使用说明书")
    _set_run_font(run, size=18, bold=True, color=BLUE)

    version = doc.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = version.add_run(f"版本：{SOFTWARE_VERSION}")
    _set_run_font(run, size=12, color=MUTED)

    generated = doc.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = generated.add_run(f"编制日期：{date.today().isoformat()}")
    _set_run_font(run, size=10.5, color=MUTED)
    doc.add_page_break()


def _add_docx_metadata(doc: Document) -> None:
    doc.add_heading("文档信息", level=1)
    rows = [
        ("软件全称", SOFTWARE_NAME),
        ("软件简称", SOFTWARE_SHORT_NAME),
        ("版本号", SOFTWARE_VERSION),
        ("使用形态", "桌面端 Web 应用"),
        ("主要用户", "新材料数据分析与集成聚类实验人员"),
        ("材料范围", "原创平台运行时代码及对应界面说明"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, (label, value) in enumerate(rows):
        cells = table.rows[row_index].cells
        cells[0].width = Inches(1.6)
        cells[1].width = Inches(4.9)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cells[0], LIGHT_FILL)
        label_run = cells[0].paragraphs[0].add_run(label)
        _set_run_font(label_run, size=10.5, bold=True, color=DARK_BLUE)
        value_run = cells[1].paragraphs[0].add_run(value)
        _set_run_font(value_run, size=10.5)
    doc.add_paragraph()


def _add_docx_content(doc: Document) -> None:
    for heading, paragraphs in SECTIONS:
        doc.add_heading(heading, level=1)
        for text in paragraphs:
            paragraph = doc.add_paragraph(text)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Pt(22)

    doc.add_heading("9. 基本操作步骤", level=1)
    steps = [
        "启动后端和前端服务，访问系统首页并确认后端连接状态正常。",
        "进入注册或登录页面完成身份认证，登录成功后进入分析工作台。",
        "在数据管理页面上传 MAT 数据集，等待结构校验和质量检查完成。",
        "在任务中心创建任务，选择数据集与算法模式并配置核心参数。",
        "启动任务后查看排队、运行进度和日志；任务完成后进入结果中心。",
        "核对指标、图表和产物清单，按需生成数据集 ZIP、结果档案或分析报告。",
        "删除数据集前先处理所有关联分析任务；退出系统时使用右上角退出按钮。",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("10. 界面截图", level=1)
    note = doc.add_paragraph(
        "以下截图来自 2026-08-10 本地开发环境的真实运行页面。截图账号未导入示例数据，数据集、任务和结果使用系统真实空状态。",
    )
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for screenshot_index, (filename, caption) in enumerate(SCREENSHOTS):
        path = SCREENSHOT_DIR / filename
        if not path.exists():
            raise FileNotFoundError(path)
        if screenshot_index > 0:
            doc.add_page_break()
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        with Image.open(path) as image:
            width, height = image.size
        max_width = 6.4
        max_height = 7.8
        scale = min(max_width / width, max_height / height) * width
        picture_width = Inches(scale)
        paragraph.add_run().add_picture(str(path), width=picture_width)
        caption_p = doc.add_paragraph(caption)
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_p.paragraph_format.space_before = Pt(8)
        caption_p.paragraph_format.keep_with_next = True
        for run in caption_p.runs:
            _set_run_font(run, size=9.5, color=MUTED)


def _write_docx(output_path: Path) -> None:
    doc = Document()
    _configure_docx_styles(doc)
    _add_docx_header_footer(doc)
    _add_docx_cover(doc)
    _add_docx_metadata(doc)
    _add_docx_content(doc)
    doc.core_properties.title = f"{SOFTWARE_NAME} 软件使用说明书"
    doc.core_properties.subject = "软件著作权登记配套材料"
    doc.core_properties.comments = "由仓库内生成脚本根据当前软件版本生成"
    doc.save(output_path)


def _register_pdf_fonts() -> None:
    pdfmetrics.registerFont(TTFont("ManualBody", "C:/Windows/Fonts/msyh.ttc", subfontIndex=0))
    pdfmetrics.registerFont(TTFont("ManualBold", "C:/Windows/Fonts/simhei.ttf"))


def _pdf_page(canvas, doc) -> None:
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D2D9E0"))
    canvas.line(22 * mm, A4[1] - 17 * mm, A4[0] - 22 * mm, A4[1] - 17 * mm)
    canvas.setFont("ManualBody", 8)
    canvas.setFillColor(colors.HexColor("#5C6873"))
    canvas.drawString(22 * mm, A4[1] - 13 * mm, f"{SOFTWARE_SHORT_NAME} | 软件使用说明书")
    canvas.drawRightString(A4[0] - 22 * mm, 12 * mm, f"第 {doc.page} 页")
    canvas.restoreState()


def _write_pdf(output_path: Path) -> None:
    _register_pdf_fonts()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ManualTitle", parent=styles["Title"], fontName="ManualBold", fontSize=24,
        leading=34, textColor=colors.HexColor("#1F4D78"), alignment=TA_CENTER, spaceAfter=12,
    )
    subtitle_style = ParagraphStyle(
        "ManualSubtitle", parent=styles["Heading1"], fontName="ManualBold", fontSize=18,
        leading=25, textColor=colors.HexColor("#2E74B5"), alignment=TA_CENTER, spaceAfter=28,
    )
    h1 = ParagraphStyle(
        "ManualH1", parent=styles["Heading1"], fontName="ManualBold", fontSize=15,
        leading=22, textColor=colors.HexColor("#2E74B5"), spaceBefore=12, spaceAfter=8,
    )
    body = ParagraphStyle(
        "ManualBody", parent=styles["BodyText"], fontName="ManualBody", fontSize=10.5,
        leading=17, alignment=TA_LEFT, firstLineIndent=21, spaceAfter=7,
    )
    small = ParagraphStyle(
        "ManualSmall", parent=body, fontSize=9, leading=14, firstLineIndent=0,
        alignment=TA_CENTER, textColor=colors.HexColor("#5C6873"),
    )
    step_style = ParagraphStyle(
        "ManualStep", parent=body, leftIndent=18, firstLineIndent=-18, bulletIndent=0,
    )

    pdf = SimpleDocTemplate(
        str(output_path), pagesize=A4, rightMargin=22 * mm, leftMargin=22 * mm,
        topMargin=22 * mm, bottomMargin=20 * mm,
        title=f"{SOFTWARE_NAME} 软件使用说明书", author=SOFTWARE_SHORT_NAME,
    )
    story = [
        Spacer(1, 50 * mm),
        Paragraph("软件著作权登记配套材料", small),
        Spacer(1, 8 * mm),
        Paragraph(SOFTWARE_NAME, title_style),
        Paragraph("软件使用说明书", subtitle_style),
        Paragraph(f"版本：{SOFTWARE_VERSION}", small),
        Paragraph(f"编制日期：{date.today().isoformat()}", small),
        PageBreak(),
        Paragraph("文档信息", h1),
    ]

    metadata = [
        ["软件全称", SOFTWARE_NAME],
        ["软件简称", SOFTWARE_SHORT_NAME],
        ["版本号", SOFTWARE_VERSION],
        ["使用形态", "桌面端 Web 应用"],
        ["主要用户", "新材料数据分析与集成聚类实验人员"],
        ["材料范围", "原创平台运行时代码及对应界面说明"],
    ]
    metadata_table = Table(metadata, colWidths=[34 * mm, 118 * mm], hAlign="LEFT")
    metadata_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "ManualBody"),
        ("FONTNAME", (0, 0), (0, -1), "ManualBold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EEF5")),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#1F4D78")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C7D0D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([metadata_table, Spacer(1, 5 * mm)])

    for heading, paragraphs in SECTIONS:
        story.append(Paragraph(heading, h1))
        story.extend(Paragraph(text, body) for text in paragraphs)

    story.append(Paragraph("9. 基本操作步骤", h1))
    for index, step in enumerate([
        "启动后端和前端服务，访问系统首页并确认后端连接状态正常。",
        "进入注册或登录页面完成身份认证，登录成功后进入分析工作台。",
        "在数据管理页面上传 MAT 数据集，等待结构校验和质量检查完成。",
        "在任务中心创建任务，选择数据集与算法模式并配置核心参数。",
        "启动任务后查看排队、运行进度和日志；任务完成后进入结果中心。",
        "核对指标、图表和产物清单，按需生成数据集 ZIP、结果档案或分析报告。",
        "删除数据集前先处理所有关联分析任务；退出系统时使用右上角退出按钮。",
    ], start=1):
        story.append(Paragraph(f"{index}. {step}", step_style))

    story.extend([
        Paragraph("10. 界面截图", h1),
        Paragraph(
            "以下截图来自 2026-08-10 本地开发环境的真实运行页面。截图账号未导入示例数据，数据集、任务和结果使用系统真实空状态。",
            body,
        ),
    ])
    for screenshot_index, (filename, caption) in enumerate(SCREENSHOTS):
        path = SCREENSHOT_DIR / filename
        with Image.open(path) as image:
            width, height = image.size
        max_width = 165 * mm
        max_height = 205 * mm
        scale = min(max_width / width, max_height / height)
        if screenshot_index > 0:
            story.append(PageBreak())
        story.append(KeepTogether([
            PdfImage(str(path), width=width * scale, height=height * scale),
            Spacer(1, 3 * mm),
            Paragraph(caption, small),
        ]))

    pdf.build(story, onFirstPage=_pdf_page, onLaterPages=_pdf_page)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    docx_path = output_dir / "新材料数据分析系统-V1.0-软件使用说明书.docx"
    pdf_path = output_dir / "新材料数据分析系统-V1.0-软件使用说明书.pdf"
    _write_docx(docx_path)
    _write_pdf(pdf_path)
    print(f"DOCX: {docx_path}")
    print(f"PDF: {pdf_path}")


if __name__ == "__main__":
    main()
